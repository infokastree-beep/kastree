"""Build a statutory DOCX from a JSON-ready statement payload.

Every string written into the document is escaped when it begins with a
formula-leading character. The builder does not calculate figures.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path

_FORMULA_LEAD = ("=", "+", "-", "@")


def escape_export_text(value: str) -> str:
    """Prefix formula-leading text so a later spreadsheet open does not run it."""
    if value.startswith(_FORMULA_LEAD):
        return "'" + value
    return value


def build_docx(payload: Mapping[str, object], dest: Path) -> None:
    """Write one DOCX. Imported after the child applies its memory cap."""
    from docx import Document

    document = Document()
    composed = payload.get("composed")
    # The composed file already prints DRAFT or FINAL in the page header.
    # A body copy sits on the cover as a second, and with the header a third.
    if not isinstance(composed, list):
        document.add_paragraph(escape_export_text(_required_text(payload, "watermark")))
    document.add_heading(
        escape_export_text(_required_text(payload, "company_name")), level=1
    )
    if isinstance(composed, list):
        _write_composed(document, payload, composed)
        document.save(str(dest))
        return
    for page in _mappings(payload.get("pages"), "pages"):
        _start_new_page(document, page.get("new_page"))
        document.add_heading(
            escape_export_text(_required_text(page, "heading")), level=2
        )
        for paragraph in _strings(page.get("paragraphs"), "paragraphs"):
            document.add_paragraph(escape_export_text(paragraph))
    _start_new_page(document, payload.get("sofp_new_page"))
    document.add_heading("Statement of financial position", level=2)
    compliance = payload.get("compliance_statement")
    if isinstance(compliance, str) and compliance:
        document.add_paragraph(escape_export_text(compliance))
    _amount_table(document, _mappings(payload.get("sofp"), "sofp"))
    _start_new_page(document, payload.get("income_new_page"))
    document.add_heading("Income statement", level=2)
    _amount_table(document, _mappings(payload.get("income"), "income"))
    _start_new_page(document, payload.get("notes_new_page"))
    document.add_heading("Notes", level=2)
    for note in _mappings(payload.get("notes"), "notes"):
        title = _required_text(note, "title")
        code = _required_text(note, "code")
        heading = code if not title else f"{code} {title}"
        document.add_heading(escape_export_text(heading), level=3)
        document.add_paragraph(escape_export_text(_required_text(note, "body")))
        lines = note.get("lines")
        if lines:
            _amount_table(document, _mappings(lines, "lines"), label_key="line")
    document.save(str(dest))


def _start_new_page(document: object, enabled: object) -> None:
    """A true flag inserts a page break. A missing flag keeps the old flow."""
    if enabled is not True:
        return
    add_page_break = getattr(document, "add_page_break", None)
    if not callable(add_page_break):
        raise TypeError("document is missing")
    add_page_break()


_NOTES_SENTENCE = "The notes form part of these financial statements."
_NOTES_FOOTER = "The notes form part of these financial statements"
_TOTAL_LABELS = frozenset(
    {
        "Total current assets",
        "Net current assets",
        "Total assets less current liabilities",
        "Net assets",
        "Total equity",
        "Gross profit",
        "Operating profit",
        "Profit before tax",
        "Profit for the financial year",
    }
)


def _write_composed(
    document: object, payload: Mapping[str, object], sections: Sequence[object]
) -> None:
    """Render the composed sections the PDF HTML already walks.

    Page breaks are the pack ``new_page`` flags already stored on each
    section. A footer change uses a Word section so the notes sentence
    appears only on the income statement, balance sheet, and notes.
    """
    add_heading = getattr(document, "add_heading", None)
    add_paragraph = getattr(document, "add_paragraph", None)
    if not callable(add_heading) or not callable(add_paragraph):
        raise TypeError("document is missing")
    page_header = payload.get("page_header")
    if not isinstance(page_header, str) or not page_header:
        page_header = _required_text(payload, "watermark")
    signature = payload.get("signature")
    if not isinstance(signature, Mapping):
        signature = {
            "date_line": "Approval date has not been recorded.",
            "names": [],
        }
    contents_anchors = _contents_anchors(sections)
    banner = _banner_texts(payload, page_header)
    notes_footer = False
    _configure_word_section(
        _word_section(document, 0),
        page_header=page_header,
        notes_footer=False,
        unlink=False,
    )
    _mark_fields_for_update(document)
    for section in sections:
        if not isinstance(section, Mapping):
            raise ValueError("composed section is missing")
        wants_footer = section.get("notes_footer") is True
        wants_break = section.get("page_break") is True
        if wants_footer != notes_footer:
            _append_word_section(
                document,
                page_break=wants_break,
                page_header=page_header,
                notes_footer=wants_footer,
            )
            notes_footer = wants_footer
        elif wants_break:
            _start_new_page(document, True)
        heading = escape_export_text(_required_text(section, "heading"))
        anchor = section.get("anchor")
        if heading.strip() not in banner:
            if isinstance(anchor, str) and anchor in contents_anchors:
                add_heading(heading, level=2)
            else:
                title = add_paragraph(heading)
                _bold_paragraph(title)
        kind = section.get("kind")
        if kind == "contents":
            labels = [
                _required_text(entry, "label")
                for entry in _mappings(section.get("entries"), "entries")
            ]
            _add_contents_field(document, labels)
            continue
        if kind == "statement":
            phrase = section.get("period_phrase")
            if isinstance(phrase, str) and phrase:
                add_paragraph(escape_export_text(phrase))
            period_note = section.get("period_note")
            if isinstance(period_note, str) and period_note:
                add_paragraph(escape_export_text(period_note))
            compliance = section.get("compliance")
            if isinstance(compliance, str) and compliance:
                add_paragraph(escape_export_text(compliance))
            rows = _mappings(section.get("rows"), "rows")
            _composed_table(
                document,
                _strings(section.get("columns"), "columns"),
                rows,
                show_notes=section.get("show_notes") is True,
            )
            if anchor == "sofp":
                add_paragraph(_NOTES_SENTENCE)
                _write_signature(document, signature)
            continue
        if kind == "notes":
            columns = _strings(section.get("columns"), "columns")
            for note in _mappings(section.get("notes"), "notes"):
                title = _required_text(note, "title").strip()
                note_heading = add_paragraph(escape_export_text(title or "Note"))
                _bold_paragraph(note_heading)
                add_paragraph(escape_export_text(_required_text(note, "body")))
                lines = _mappings(note.get("lines"), "lines")
                if lines:
                    _composed_table(document, columns, lines)
                fa_rows = note.get("fa_rows")
                if isinstance(fa_rows, list) and fa_rows:
                    _fa_table(document, _mappings(fa_rows, "fa_rows"))
            continue
        for paragraph in _strings(section.get("paragraphs"), "paragraphs"):
            shown = escape_export_text(paragraph)
            if shown.strip() in banner:
                continue
            add_paragraph(shown)


def _banner_texts(payload: Mapping[str, object], page_header: str) -> frozenset[str]:
    """Labels that belong in the header, not again in the cover body."""
    texts = {page_header.strip()}
    watermark = payload.get("watermark")
    if isinstance(watermark, str) and watermark.strip():
        texts.add(watermark.strip())
    return frozenset(texts)


def _contents_anchors(sections: Sequence[object]) -> frozenset[str]:
    anchors: set[str] = set()
    for section in sections:
        if not isinstance(section, Mapping) or section.get("kind") != "contents":
            continue
        for entry in _mappings(section.get("entries"), "entries"):
            anchor = entry.get("anchor")
            if isinstance(anchor, str):
                anchors.add(anchor)
    return frozenset(anchors)


def _write_signature(document: object, signature: Mapping[str, object]) -> None:
    add_paragraph = getattr(document, "add_paragraph", None)
    if not callable(add_paragraph):
        raise TypeError("document is missing")
    add_paragraph("Approved by the board and signed on its behalf by")
    date_line = signature.get("date_line")
    if not isinstance(date_line, str) or not date_line:
        date_line = "Approval date has not been recorded."
    add_paragraph(escape_export_text(date_line))
    names = signature.get("names")
    printed = (
        [name.strip() for name in names if isinstance(name, str) and name.strip()]
        if isinstance(names, list)
        else []
    )
    if not printed:
        add_paragraph("A signatory has not been recorded.")
        add_paragraph("______________________________")
        add_paragraph("Director")
        return
    for name in printed:
        add_paragraph("______________________________")
        add_paragraph(f"{escape_export_text(name)}\nDirector")


def _composed_table(
    document: object,
    columns: Sequence[str],
    rows: Sequence[Mapping[str, object]],
    *,
    show_notes: bool = False,
) -> None:
    if not rows:
        return
    add_table = getattr(document, "add_table", None)
    if not callable(add_table):
        raise TypeError("document is missing")
    headings = ["Line", *columns]
    note_at: int | None = None
    if show_notes:
        headings = ["Line", "Notes", *columns]
        note_at = 1
    amount_at = set(range(1 if note_at is None else 2, len(headings)))
    table = add_table(rows=1 + len(rows), cols=len(headings))
    header = table.rows[0].cells
    for index, column in enumerate(headings):
        header[index].text = escape_export_text(column)
    total_rows: set[int] = set()
    for index, row in enumerate(rows, start=1):
        cells = table.rows[index].cells
        label = escape_export_text(_required_text(row, "label"))
        cells[0].text = label
        if note_at is not None:
            note = row.get("note")
            cells[note_at].text = escape_export_text(
                note if isinstance(note, str) else ""
            )
        amounts = row.get("amounts")
        if not isinstance(amounts, list):
            raise ValueError("composed amounts are missing")
        for offset, amount in enumerate(amounts):
            if not isinstance(amount, str):
                raise ValueError("composed amount is missing")
            column_index = (2 if note_at is not None else 1) + offset
            if column_index < len(headings):
                cells[column_index].text = escape_export_text(amount)
        if _is_total(label):
            total_rows.add(index)
    _format_amount_table(
        table,
        amount_columns=amount_at,
        total_rows=total_rows,
        column_count=len(headings),
        notes_column=note_at is not None,
    )


def _fa_table(document: object, rows: Sequence[Mapping[str, object]]) -> None:
    if not rows:
        return
    add_table = getattr(document, "add_table", None)
    if not callable(add_table):
        raise TypeError("document is missing")
    table = add_table(rows=1 + len(rows), cols=2)
    header = table.rows[0].cells
    header[0].text = "Class"
    header[1].text = "NBV"
    for index, row in enumerate(rows, start=1):
        cells = table.rows[index].cells
        cells[0].text = escape_export_text(_required_text(row, "asset_class"))
        cells[1].text = escape_export_text(_required_text(row, "nbv_close"))
    _format_amount_table(
        table,
        amount_columns={1},
        total_rows=set(),
        column_count=2,
        notes_column=False,
    )


def _is_total(label: str) -> bool:
    plain = label[1:] if label.startswith("'") else label
    return plain in _TOTAL_LABELS or plain.lower().startswith("total ")


def _bold_paragraph(paragraph: object) -> None:
    runs = getattr(paragraph, "runs", None)
    if not isinstance(runs, list):
        return
    for run in runs:
        if hasattr(run, "bold"):
            run.bold = True


def _word_section(document: object, index: int) -> object:
    sections = getattr(document, "sections", None)
    if sections is None:
        raise TypeError("document is missing")
    return sections[index]


def _append_word_section(
    document: object,
    *,
    page_break: bool,
    page_header: str,
    notes_footer: bool,
) -> None:
    from docx.enum.section import WD_SECTION

    add_section = getattr(document, "add_section", None)
    if not callable(add_section):
        raise TypeError("document is missing")
    kind = WD_SECTION.NEW_PAGE if page_break else WD_SECTION.CONTINUOUS
    section = add_section(kind)
    _configure_word_section(
        section,
        page_header=page_header,
        notes_footer=notes_footer,
        unlink=True,
    )


def _configure_word_section(
    section: object, *, page_header: str, notes_footer: bool, unlink: bool
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Mm, Pt, RGBColor

    if getattr(section, "page_width", None) is not None:
        setattr(section, "page_width", Mm(210))
        setattr(section, "page_height", Mm(297))
        setattr(section, "left_margin", Mm(14))
        setattr(section, "right_margin", Mm(14))
        setattr(section, "top_margin", Mm(16))
        setattr(section, "bottom_margin", Mm(22))
    header = getattr(section, "header", None)
    footer = getattr(section, "footer", None)
    if header is None or footer is None:
        raise TypeError("document is missing")
    if unlink:
        header.is_linked_to_previous = False
        footer.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = ""
    run = header_paragraph.add_run(escape_export_text(page_header))
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.text = ""
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(
        Mm(182), WD_TAB_ALIGNMENT.RIGHT
    )
    if notes_footer:
        footer_run = footer_paragraph.add_run(_NOTES_FOOTER)
        _style_footer_run(footer_run)
    footer_paragraph.add_run("\t")
    page_run = footer_paragraph.add_run()
    _style_footer_run(page_run)
    _append_field(page_run._r, " PAGE ")
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _style_footer_run(run: object) -> None:
    from docx.shared import Pt, RGBColor

    font = getattr(run, "font", None)
    if font is None:
        return
    font.size = Pt(8)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def _add_contents_field(document: object, labels: Sequence[str]) -> None:
    """A TOC field. Word shows the labels until the reader updates fields."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    add_paragraph = getattr(document, "add_paragraph", None)
    if not callable(add_paragraph):
        raise TypeError("document is missing")
    paragraph = add_paragraph()
    run = paragraph.add_run()
    element = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    element.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "2-2" \h \z \u '
    element.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    element.append(separate)
    if not labels:
        shown = OxmlElement("w:t")
        shown.text = "Contents will list each section after fields are updated."
        element.append(shown)
    for index, label in enumerate(labels):
        if index:
            element.append(OxmlElement("w:br"))
        shown = OxmlElement("w:t")
        shown.set(qn("xml:space"), "preserve")
        shown.text = escape_export_text(label)
        element.append(shown)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    element.append(end)


def _append_field(run_element: object, instruction: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    append = getattr(run_element, "append", None)
    if not callable(append):
        raise TypeError("document is missing")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    append(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    append(end)


def _mark_fields_for_update(document: object) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = getattr(getattr(document, "settings", None), "element", None)
    if settings is None:
        return
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _format_amount_table(
    table: object,
    *,
    amount_columns: set[int],
    total_rows: set[int],
    column_count: int,
    notes_column: bool,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm

    tbl = getattr(table, "_tbl", None)
    if tbl is None:
        raise TypeError("document is missing")
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    usable = 182
    notes_width = 12
    amount_width = 22
    if notes_column:
        other = notes_width + amount_width * max(column_count - 2, 0)
    else:
        other = amount_width * max(column_count - 1, 0)
    label_width = max(usable - other, 40)
    widths = [label_width]
    if notes_column:
        widths.append(notes_width)
    widths.extend([amount_width] * (column_count - len(widths)))
    _set_table_widths(tbl, widths)
    rows = getattr(table, "rows", None)
    if rows is None:
        raise TypeError("document is missing")
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Mm(widths[column_index])
            _cell_bottom_rule(cell)
            align_right = column_index in amount_columns
            bold = row_index in total_rows
            for paragraph in cell.paragraphs:
                if align_right:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if bold:
                    for run in paragraph.runs:
                        run.bold = True


def _set_table_widths(tbl: object, widths_mm: list[int]) -> None:
    """Write the grid Word uses. Cell widths alone leave the columns equal."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm

    find = getattr(tbl, "find", None)
    tbl_pr = getattr(tbl, "tblPr", None)
    if not callable(find) or tbl_pr is None:
        raise TypeError("document is missing")
    twips = [int(Mm(width).twips) for width in widths_mm]
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(twips)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for width in twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)


def _cell_bottom_rule(cell: object) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = getattr(cell, "_tc", None)
    if tc is None:
        return
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "CCCCCC")
    borders.append(bottom)
    for edge in ("top", "left", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def _amount_table(
    document: object,
    rows: Sequence[Mapping[str, object]],
    *,
    label_key: str = "label",
) -> None:
    add_table = getattr(document, "add_table", None)
    if not callable(add_table):
        raise TypeError("document is missing")
    table = add_table(rows=1 + len(rows), cols=3)
    headers = table.rows[0].cells
    headers[0].text = "Line"
    headers[1].text = "Current"
    headers[2].text = "Prior"
    for index, row in enumerate(rows, start=1):
        cells = table.rows[index].cells
        cells[0].text = escape_export_text(_required_text(row, label_key))
        cells[1].text = escape_export_text(_required_text(row, "current"))
        cells[2].text = escape_export_text(_required_text(row, "prior"))


def _required_text(payload: Mapping[str, object], key: str) -> str:
    value = payload.get(key)
    if not isinstance(value, str):
        raise ValueError(f"{key} is missing")
    return value


def _mappings(value: object, key: str) -> tuple[Mapping[str, object], ...]:
    if not isinstance(value, list):
        raise ValueError(f"{key} is missing")
    rows: list[Mapping[str, object]] = []
    for item in value:
        if not isinstance(item, dict):
            raise ValueError(f"{key} is missing")
        rows.append(item)
    return tuple(rows)


def _strings(value: object, key: str) -> tuple[str, ...]:
    if not isinstance(value, list):
        raise ValueError(f"{key} is missing")
    texts: list[str] = []
    for item in value:
        if not isinstance(item, str):
            raise ValueError(f"{key} is missing")
        texts.append(item)
    return tuple(texts)
